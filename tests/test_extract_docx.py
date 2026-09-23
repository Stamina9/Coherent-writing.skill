import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import struct
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/coherent-academic-writing/scripts/extract_docx.py"
SPEC = importlib.util.spec_from_file_location("extract_docx", SCRIPT)
extractor = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(extractor)


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.source = self.directory / "论文.docx"

    def tearDown(self):
        self.temp.cleanup()

    def save(self, doc):
        doc.save(self.source)
        result = extractor.extract_docx(self.source)
        self.assert_contract(result)
        return result

    def assert_contract(self, result):
        """Validate the public JSON 1.0 contract, not the extractor's internals."""
        self.assertEqual(result["schema_version"], "1.0")
        self.assertIsInstance(result["source"], str)
        self.assertTrue(Path(result["source"]).is_absolute())
        self.assertRegex(result["source_sha256"], r"^[0-9a-f]{64}$")
        self.assertIsInstance(result["blocks"], list)
        self.assertIsInstance(result["warnings"], list)
        ids = set()
        paragraphs = tables = 0
        for block in result["blocks"]:
            self.assertIsInstance(block["id"], str)
            self.assertNotIn(block["id"], ids)
            ids.add(block["id"])
            if block["kind"] == "paragraph":
                paragraphs += 1
                self.assertIsInstance(block["text"], str)
                self.assertTrue(block["style"] is None or isinstance(block["style"], str))
                level = block["heading_level"]
                self.assertTrue(level is None or type(level) is int and 1 <= level <= 9)
                self.assertIs(type(block["caption_candidate"]), bool)
            else:
                self.assertEqual(block["kind"], "table")
                tables += 1
                self.assertIs(type(block["rows"]), int)
                self.assertGreaterEqual(block["rows"], 0)
        for warning in result["warnings"]:
            for key in ("code", "location", "detail"):
                self.assertIsInstance(warning[key], str)
        coverage = result["coverage"]
        self.assertIsInstance(coverage["scope"], str)
        self.assertIn(coverage["status"], ("partial", "text-only"))
        self.assertEqual(coverage["paragraph_count"], paragraphs)
        self.assertEqual(coverage["table_count"], tables)
        for key in ("excluded_parts", "limitations"):
            self.assertIsInstance(coverage[key], list)
            self.assertTrue(all(isinstance(v, str) for v in coverage[key]))

    def codes(self, result):
        return {w["code"] for w in result["warnings"]}

    def test_full_text_order_tables_and_source_unchanged(self):
        doc = Document()
        doc.add_heading("发现与解释", 1)
        long_text = "仅在本样本中成立，负面结果仍需报告。" * 80
        doc.add_paragraph(long_text)
        table = doc.add_table(rows=2, cols=2)
        for cell, value in zip([c for r in table.rows for c in r.cells], ["组别", "效应", "干预组", "−2.1 (95% CI −3.2 至 −1.0)"]):
            cell.text = value
        doc.add_paragraph("表1 全部结果", style="Caption")
        doc.add_paragraph("")
        doc.add_paragraph("尾段\t含制表符\n以及换行")
        doc.save(self.source)
        before = self.source.read_bytes()
        result = extractor.extract_docx(self.source)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(result["source_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(result["blocks"][1]["text"], long_text)
        self.assertEqual(result["blocks"][2]["id"], "T0")
        self.assertEqual(result["blocks"][6]["text"], "−2.1 (95% CI −3.2 至 −1.0)")
        self.assertEqual(result["blocks"][7]["id"], "P2")
        self.assertTrue(result["blocks"][7]["caption_candidate"])
        self.assertEqual(result["blocks"][-1]["text"], "尾段\t含制表符\n以及换行")
        self.assertEqual(result, extractor.extract_docx(self.source))

    def test_nonstandard_heading_and_inherited_outline(self):
        doc = Document()
        style = doc.styles.add_style("Local section", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 2"]
        doc.add_paragraph("命题与证明", style=style)
        normal = doc.add_paragraph("自定义章节")
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        normal._p.get_or_add_pPr().append(outline)
        result = self.save(doc)
        self.assertEqual([b["heading_level"] for b in result["blocks"]], [2, 1])
        self.assertNotIn("NO_HEADINGS", self.codes(result))

    def test_no_heading_preserves_candidate_text(self):
        doc = Document()
        doc.add_paragraph("3 发现与解释")
        doc.add_paragraph("仅适用于已测试条件。")
        result = self.save(doc)
        self.assertIn("NO_HEADINGS", self.codes(result))
        self.assertTrue(all(b["heading_level"] is None for b in result["blocks"]))
        self.assertEqual(result["blocks"][0]["text"], "3 发现与解释")

    def test_nested_and_merged_tables_keep_unique_content(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "合并证据"
        merged.add_table(rows=1, cols=1).cell(0, 0).text = "嵌套反例"
        result = self.save(doc)
        texts = [b.get("text") for b in result["blocks"]]
        self.assertEqual(texts.count("合并证据"), 1)
        self.assertEqual(texts.count("嵌套反例"), 1)
        self.assertEqual(result["coverage"]["table_count"], 2)
        self.assertIn("IRREGULAR_TABLE", self.codes(result))
        ids = [b["id"] for b in result["blocks"]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_hyperlink_visible_text(self):
        doc = Document()
        para = doc.add_paragraph("证据见")
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "附录 A"
        run.append(text)
        link.append(run)
        para._p.append(link)
        self.assertEqual(self.save(doc)["blocks"][0]["text"], "证据见附录 A")

    def test_incomplete_objects_are_observable(self):
        doc = Document()
        p = doc.add_paragraph("可读取正文")
        p._p.append(OxmlElement("w:ins"))
        p._p.append(OxmlElement("m:oMath"))
        p._p.append(OxmlElement("w:drawing"))
        p._p.append(OxmlElement("w:txbxContent"))
        p._p.append(OxmlElement("w:fldSimple"))
        doc.element.body.append(OxmlElement("w:sdt"))
        doc.sections[0].header.paragraphs[0].text = "页眉证据"
        result = self.save(doc)
        self.assertTrue({"DRAWING_NOT_EXTRACTED", "TEXTBOX_NOT_EXTRACTED", "MATH_NOT_EXTRACTED",
                         "REVISIONS_NOT_RESOLVED", "FIELD_NOT_EVALUATED", "PART_NOT_EXTRACTED",
                         "CONTENT_CONTROL_NOT_EXTRACTED"}.issubset(self.codes(result)))
        self.assertEqual(result["coverage"]["status"], "partial")
        self.assertIn("word/header1.xml", result["coverage"]["excluded_parts"])

    def test_bad_docx_and_encrypted_or_legacy_errors(self):
        for payload, code in [(b"bad zip", "INVALID_DOCX"), (extractor.OLE_MAGIC + b"test", "ENCRYPTED_OR_LEGACY")]:
            with self.subTest(code=code):
                self.source.write_bytes(payload)
                error = io.StringIO()
                with contextlib.redirect_stderr(error):
                    self.assertEqual(extractor.main([str(self.source)]), 4)
                self.assertIn(code, error.getvalue())
                self.assertEqual(self.source.read_bytes(), payload)

    def test_malformed_xml_and_non_docx_zip(self):
        for include_xml in [False, True]:
            with self.subTest(include_xml=include_xml):
                with zipfile.ZipFile(self.source, "w") as package:
                    package.writestr("[Content_Types].xml", "<invalid>")
                    if include_xml:
                        package.writestr("word/document.xml", "<invalid>")
                with self.assertRaises(extractor.ExtractionError) as caught:
                    extractor.extract_docx(self.source)
                self.assertEqual(caught.exception.code, "INVALID_DOCX")

    def test_empty_document_stops(self):
        with self.assertRaises(extractor.ExtractionError) as caught:
            self.save(Document())
        self.assertEqual(caught.exception.code, "EMPTY_TEXT")

    def test_resource_limit(self):
        doc = Document()
        doc.add_paragraph("正文")
        doc.save(self.source)
        with patch.object(extractor, "MAX_EXPANDED", 1):
            with self.assertRaises(extractor.ExtractionError) as caught:
                extractor.extract_docx(self.source)
        self.assertEqual(caught.exception.code, "PACKAGE_LIMIT")

    def test_missing_dependency_in_isolated_interpreter(self):
        # -I -S really removes site-packages; no fake fixed error-string implementation.
        result = subprocess.run([sys.executable, "-I", "-S", str(SCRIPT), str(self.source)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 3)
        self.assertIn("MISSING_DEPENDENCY", result.stderr)
        self.assertEqual(result.stdout, "")
        self.assertNotIn("Traceback", result.stderr)

    def test_cli_unique_outputs_and_failure_cleanup(self):
        doc = Document()
        doc.add_paragraph("事实没有改变。")
        doc.save(self.source)
        paths = []
        for _ in range(2):
            result = subprocess.run([sys.executable, "-X", "utf8", str(SCRIPT), str(self.source), "--output-dir", str(self.directory)], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 0, result.stderr)
            paths.append(Path(result.stdout.strip()))
        self.assertNotEqual(*paths)
        self.assert_contract(json.loads(paths[0].read_text(encoding="utf-8")))
        self.assertEqual(json.loads(paths[0].read_text(encoding="utf-8")), json.loads(paths[1].read_text(encoding="utf-8")))
        before = set(self.directory.iterdir())
        with patch.object(extractor.json, "dump", side_effect=OSError("disk full")):
            with self.assertRaises(extractor.ExtractionError):
                extractor.write_result({}, self.directory)
        self.assertEqual(set(self.directory.iterdir()), before)

    def test_io_errors_are_actionable(self):
        with self.assertRaises(extractor.ExtractionError) as caught:
            extractor.extract_docx(self.directory / "missing.docx")
        self.assertEqual(caught.exception.exit_code, 5)
        self.source.write_bytes(b"occupied")
        with self.assertRaises(extractor.ExtractionError) as caught:
            extractor.write_result({}, self.source)
        self.assertEqual(caught.exception.code, "OUTPUT_IO")
        self.assertEqual(self.source.read_bytes(), b"occupied")

    def test_corrupt_deflate_returns_diagnostic_without_output(self):
        doc = Document()
        doc.add_paragraph("text")
        doc.save(self.source)
        raw = bytearray(self.source.read_bytes())
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            offset = package.getinfo("word/document.xml").header_offset
        name_len, extra_len = struct.unpack_from("<HH", raw, offset + 26)
        # BTYPE=3 is an invalid DEFLATE block, while the ZIP directory remains valid.
        raw[offset + 30 + name_len + extra_len] = 7
        self.source.write_bytes(raw)
        output = self.directory / "output"
        result = subprocess.run([sys.executable, str(SCRIPT), str(self.source), "--output-dir", str(output)], capture_output=True)
        self.assertEqual(result.returncode, 4, result.stderr)
        self.assertIn(b"INVALID_DOCX", result.stderr)
        self.assertNotIn(b"Traceback", result.stderr)
        self.assertEqual(result.stdout, b"")
        self.assertFalse(output.exists())
        self.assertEqual(self.source.read_bytes(), raw)

    def test_nonstandard_zip_compression_is_rejected_cleanly(self):
        with zipfile.ZipFile(self.source, "w", compression=zipfile.ZIP_BZIP2) as package:
            package.writestr("[Content_Types].xml", "<Types/>")
            package.writestr("word/document.xml", "<document/>")
        with self.assertRaises(extractor.ExtractionError) as caught:
            extractor.extract_docx(self.source)
        self.assertEqual(caught.exception.code, "INVALID_DOCX")

    def test_deleted_table_row_requires_revision_review(self):
        doc = Document()
        doc.add_heading("Results", 1)
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "No improvement"
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:del"))
        result = self.save(doc)
        self.assertIn("REVISIONS_NOT_RESOLVED", self.codes(result))
        self.assertEqual(result["coverage"]["status"], "partial")
        self.assertTrue(any(b.get("text") == "No improvement" for b in result["blocks"]))

    def test_unextracted_inline_symbol_warns_instead_of_silent_text(self):
        doc = Document()
        doc.add_heading("Results", 1)
        run = doc.add_paragraph("Effect: ").add_run()
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:font"), "Symbol")
        symbol.set(qn("w:char"), "F02D")
        run._r.append(symbol)
        run.add_text("2")
        result = self.save(doc)
        self.assertIn("INLINE_CONTENT_NOT_EXTRACTED", self.codes(result))
        self.assertEqual(result["coverage"]["status"], "partial")

    def test_cli_path_is_utf8_even_with_legacy_stdout_encoding(self):
        doc = Document()
        doc.add_paragraph("text")
        doc.save(self.source)
        output = self.directory / "中文输出"
        env = dict(os.environ, PYTHONIOENCODING="ascii")
        result = subprocess.run([sys.executable, str(SCRIPT), str(self.source), "--output-dir", str(output)], capture_output=True, env=env)
        self.assertEqual(result.returncode, 0, result.stderr)
        path = Path(result.stdout.decode("utf-8").strip())
        self.assertTrue(path.is_file())
        self.assertEqual(path.parent, output.resolve())


if __name__ == "__main__":
    unittest.main()
